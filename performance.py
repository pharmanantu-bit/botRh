"""Module PERFORMANCE (Blueprint Flask) — rentabilité des collaborateurs.

L'admin importe chaque mois l'export OSPharm « recap mois collaborateur »
(.xlsx : 1 colonne par vendeur, blocs période × indicateur). Le module croise :
  - ce que le collaborateur RAPPORTE : CA TTC reconstitué (nb ventes × panier
    moyen, global / ordonnance / hors ordonnance), volumes et CA HT par famille ;
  - ce qu'il COÛTE : brut mensuel (+ prime variable) saisi par l'admin —
    chiffré au repos comme l'IBAN — × coefficient de charges patronales ;
  - une MARGE ESTIMÉE via des taux de marge configurables (ordonnance vs hors
    ordonnance), l'export ne contenant pas la marge réelle ;
  - les OBJECTIFS d'entretien individuel (indicateur, cible, échéance) avec
    l'écart réalisé/objectif mis à jour à chaque import.

Données : performance.json (gitignoré, rémunérations chiffrées Fernet).
Les cumuls « 202509:202608 » du fichier sont ignorés : on ne stocke que les
mois simples et on recalcule les cumuls nous-mêmes.
"""
import os
import re
import unicodedata
import uuid
from datetime import datetime

from flask import (Blueprint, request, render_template, redirect, url_for,
                   session, abort)

from app import (_lire_json, _ecrire_json, _nombre_fr, BASE_DIR, MOIS_FR,
                 charger_employes, charger_profils, collaborateur_actif)
import crypto_rh

bp = Blueprint("performance", __name__)

PERF_FILE = os.path.join(BASE_DIR, "performance.json")

# Intitulés OSPharm -> clés internes (comparaison insensible casse/accents).
INDICATEURS = {
    "nombre ventes global": "nb_global",
    "nombre ventes ordonnance": "nb_ordo",
    "nombre ventes hors ordonnance": "nb_hors",
    "panier ttc moyen global": "panier_global",
    "panier ttc moyen ordonnance": "panier_ordo",
    "panier ttc moyen hors ordonnance": "panier_hors",
    "taux de substitution": "taux_subst",
    "volume vendu phytotherapie": "vol_phyto",
    "volume vendu dietetique": "vol_diet",
    "volume vendu bebe": "vol_bebe",
    "volume vendu parapharmacie": "vol_para",
    "chiffre d'affaires ht phytotherapie": "ca_ht_phyto",
    "chiffre d'affaires ht dietetique": "ca_ht_diet",
    "chiffre d'affaires ht bebe": "ca_ht_bebe",
    "chiffre d'affaires ht parapharmacie": "ca_ht_para",
}

# Indicateurs proposables comme OBJECTIF d'entretien (clé -> libellé + unité).
INDICATEURS_OBJECTIF = {
    "ca_ttc_global": ("CA TTC du mois (estimé)", "€"),
    "ca_ttc_hors": ("CA TTC hors ordonnance du mois (estimé)", "€"),
    "nb_global": ("Nombre de ventes du mois", "ventes"),
    "panier_global": ("Panier moyen TTC global", "€"),
    "panier_hors": ("Panier moyen TTC hors ordonnance", "€"),
    "taux_subst": ("Taux de substitution", "%"),
    "vol_para": ("Volumes parapharmacie du mois", "unités"),
    "vol_diet": ("Volumes diététique du mois", "unités"),
    "rentabilite": ("Rentabilité du mois (marge estimée − coût)", "€"),
}

OPTIONS_DEFAUT = {"coef_charges": 1.42, "taux_marge_ordo": 8.0, "taux_marge_hors": 30.0}


def charger_perf():
    d = _lire_json(PERF_FILE)
    d.setdefault("mois", {})
    d.setdefault("vendeurs", {})
    d.setdefault("mapping", {})
    d.setdefault("remunerations", {})
    d.setdefault("objectifs", {})
    d.setdefault("options", dict(OPTIONS_DEFAUT))
    return d


def sauvegarder_perf(d):
    _ecrire_json(PERF_FILE, d)


def _normaliser(txt):
    txt = unicodedata.normalize("NFKD", str(txt or "").strip().lower())
    return "".join(c for c in txt if not unicodedata.combining(c))


def titre_periode(p):
    """'202608' -> 'Août 2026'."""
    try:
        return f"{MOIS_FR[int(p[4:6])]} {p[:4]}"
    except (KeyError, ValueError, IndexError):
        return p


# ------------------------------------------------------------------- Import

def importer_xlsx(chemin_ou_flux):
    """Parse l'export OSPharm. Renvoie (vendeurs {code: prénom},
    mois {'AAAAMM': {code: {cle: valeur}}}, erreurs [str])."""
    from openpyxl import load_workbook
    erreurs = []
    try:
        wb = load_workbook(chemin_ou_flux, data_only=True, read_only=True)
    except Exception as e:
        return {}, {}, [f"Fichier illisible : {e}"]
    ws = wb.active
    lignes = ws.iter_rows(values_only=True)
    try:
        l_prenoms = next(lignes)
        l_codes = next(lignes)
    except StopIteration:
        return {}, {}, ["Fichier vide."]
    if _normaliser(l_codes[0]) != "periode" or _normaliser(l_codes[1]) != "indicateur":
        return {}, {}, ["Format inattendu : colonnes PERIODE / INDICATEUR introuvables "
                        "(est-ce bien l'export OSPharm « recap mois collaborateur » ?)."]
    vendeurs, colonnes = {}, []   # colonnes : [(index, code)]
    for i in range(2, len(l_codes)):
        code = str(l_codes[i] or "").strip()
        if code:
            vendeurs[code] = str(l_prenoms[i] or "").strip() if i < len(l_prenoms) else ""
            colonnes.append((i, code))
    mois, inconnus = {}, set()
    for row in lignes:
        periode = str(row[0] or "").strip()
        if not re.fullmatch(r"\d{6}", periode):
            continue   # cumuls '202509:202608' et lignes vides : ignorés
        cle = INDICATEURS.get(_normaliser(row[1]))
        if not cle:
            if row[1]:
                inconnus.add(str(row[1]).strip())
            continue
        bloc = mois.setdefault(periode, {})
        for i, code in colonnes:
            v = row[i] if i < len(row) else None
            if v is None or v == "":
                continue
            try:
                bloc.setdefault(code, {})[cle] = round(float(v), 4)
            except (TypeError, ValueError):
                pass
    for x in sorted(inconnus):
        erreurs.append(f"Indicateur non reconnu (ignoré) : « {x} »")
    if not mois:
        erreurs.append("Aucun mois de données trouvé dans le fichier.")
    return vendeurs, mois, erreurs


def proposer_mapping(vendeurs, mapping, employes):
    """Complète mapping {code: email} par correspondance sur le prénom.
    Ne touche pas aux associations déjà faites (y compris '' = ignorer)."""
    par_prenom = {}
    for e in employes:
        par_prenom.setdefault(_normaliser(e.get("prenom")), e.get("email", ""))
    for code, prenom in vendeurs.items():
        if code not in mapping:
            mapping[code] = par_prenom.get(_normaliser(prenom), "")
    return mapping


def remuneration_effective(email, perf, profils):
    """Rémunération à utiliser pour un collaborateur : la saisie admin si elle
    existe, sinon repli sur le salaire brut de la PROMESSE D'EMBAUCHE stockée
    sur la fiche (le salaire des contrats n'est volontairement pas conservé —
    minimisation). Renvoie (rem, source) avec source 'saisie' / 'promesse' / None."""
    rem = perf["remunerations"].get(email)
    if rem:
        return rem, "saisie"
    brut = _nombre_fr((profils.get(email, {}).get("promesse") or {}).get("salaire_brut"))
    if brut:
        return {"brut": crypto_rh.chiffrer(str(brut))}, "promesse"
    return {}, None


# ------------------------------------------------------------------ Calculs

def _cout(rem, options):
    """Coût employeur mensuel : brut × coef charges + variable. None si brut absent."""
    brut = _dechiffrer_nombre(rem.get("brut"))
    if brut is None:
        return None
    variable = _dechiffrer_nombre(rem.get("variable")) or 0.0
    return brut * float(options.get("coef_charges", 1.42)) + variable


def _dechiffrer_nombre(valeur):
    if valeur in (None, ""):
        return None
    txt = crypto_rh.dechiffrer(valeur)
    try:
        return float(str(txt).replace(",", "."))
    except (TypeError, ValueError):
        return None   # clé de chiffrement absente ('🔒') ou saisie invalide


def indicateurs_derives(d):
    """Complète un bloc d'indicateurs avec les CA TTC reconstitués."""
    r = dict(d)
    for suffixe in ("global", "ordo", "hors"):
        nb, panier = r.get(f"nb_{suffixe}"), r.get(f"panier_{suffixe}")
        r[f"ca_ttc_{suffixe}"] = round(nb * panier, 2) if nb is not None and panier is not None else None
    return r


def bilan_collaborateur(d, rem, options):
    """d = indicateurs d'UN mois (ou sommés) -> ajoute marge estimée, coût,
    rentabilité et ratio. Marge = CA ordo × taux_ordo + CA hors ordo × taux_hors."""
    r = indicateurs_derives(d)
    ca_ordo, ca_hors = r.get("ca_ttc_ordo"), r.get("ca_ttc_hors")
    if ca_ordo is not None or ca_hors is not None:
        r["marge_estimee"] = round((ca_ordo or 0) * options.get("taux_marge_ordo", 8.0) / 100
                                   + (ca_hors or 0) * options.get("taux_marge_hors", 30.0) / 100, 2)
    else:
        r["marge_estimee"] = None
    cout = _cout(rem, options)
    r["cout"] = round(cout, 2) if cout is not None else None
    r["rentabilite"] = (round(r["marge_estimee"] - cout, 2)
                        if r["marge_estimee"] is not None and cout is not None else None)
    r["ratio"] = (round(r["ca_ttc_global"] / cout, 2)
                  if r.get("ca_ttc_global") and cout else None)
    return r


def sommer_mois(blocs):
    """Somme des indicateurs additifs sur plusieurs mois + paniers recalculés.
    (Un panier moyen ne s'additionne pas : on le déduit des CA/nb sommés.)"""
    total = {}
    for b in blocs:
        for k, v in b.items():
            if k.startswith(("nb_", "vol_", "ca_ht_")) and v is not None:
                total[k] = round(total.get(k, 0) + v, 2)
    # CA TTC sommés depuis les mois (nb × panier de CHAQUE mois, puis somme)
    for suffixe in ("global", "ordo", "hors"):
        cas = [b.get(f"nb_{suffixe}", 0) * b.get(f"panier_{suffixe}", 0)
               for b in blocs
               if b.get(f"nb_{suffixe}") is not None and b.get(f"panier_{suffixe}") is not None]
        if cas:
            total[f"ca_ttc_{suffixe}"] = round(sum(cas), 2)
            nb = total.get(f"nb_{suffixe}")
            if nb:
                total[f"panier_{suffixe}"] = round(total[f"ca_ttc_{suffixe}"] / nb, 2)
    return total


def valeur_indicateur(bilan, cle):
    """Valeur d'un indicateur d'objectif dans un bilan (taux en %)."""
    v = bilan.get(cle)
    if v is None:
        return None
    return round(v * 100, 1) if cle == "taux_subst" else round(v, 2)


def suivi_objectif(obj, perf, options):
    """Complète un objectif avec sa dernière valeur mesurée et son avancement.
    Mesure = mois de l'échéance s'il est importé, sinon dernier mois importé."""
    email = obj["email"]
    mois_dispo = sorted(perf["mois"])
    if not mois_dispo:
        return {**obj, "valeur": None, "mois_mesure": None, "avancement": None, "atteint": None}
    echeance = (obj.get("echeance") or "").replace("-", "")
    mois_mesure = echeance if echeance in perf["mois"] else mois_dispo[-1]
    code = next((c for c, e in perf["mapping"].items() if e == email), None)
    bloc = perf["mois"][mois_mesure].get(code or "", {})
    rem, _ = remuneration_effective(email, perf, charger_profils())
    bilan = bilan_collaborateur(bloc, rem, options)
    valeur = valeur_indicateur(bilan, obj.get("indicateur"))
    cible = obj.get("cible")
    avancement = round(100 * valeur / cible, 1) if valeur is not None and cible else None
    return {**obj, "valeur": valeur, "mois_mesure": mois_mesure,
            "avancement": avancement,
            "atteint": (valeur is not None and cible is not None and valeur >= cible)}


# -------------------------------------------------------------------- Pages

@bp.route("/admin/performance")
def admin_performance():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    perf = charger_perf()
    options = perf["options"]
    profils = charger_profils()
    # Collaborateurs ACTIFS seulement : un archivé (parti, dossier conservé)
    # est ignoré partout — tableau, cumul, correspondances, fiches.
    employes = [e for e in charger_employes()
                if collaborateur_actif(profils.get(e["email"], {}))]
    emails_noms = {e["email"]: f"{e['prenom']} {e.get('nom', '')}".strip() for e in employes}
    mois_dispo = sorted(perf["mois"], reverse=True)
    mois_sel = request.args.get("mois", "")
    if mois_sel not in perf["mois"]:
        mois_sel = mois_dispo[0] if mois_dispo else ""

    collab = request.args.get("collab", "")
    if collab and collab in emails_noms:
        return _page_collaborateur(perf, options, collab, emails_noms, mois_dispo)

    lignes, sans_collab = [], []
    if mois_sel:
        for code, bloc in sorted(perf["mois"][mois_sel].items()):
            email = perf["mapping"].get(code, "")
            if email and email not in emails_noms:
                continue   # associé à un collaborateur parti (archivé) : ignoré
            rem, rem_source = remuneration_effective(email, perf, profils)
            b = bilan_collaborateur(bloc, rem, options)
            objs = [suivi_objectif({**o, "email": email}, perf, options)
                    for o in perf["objectifs"].get(email, []) if not o.get("clos")]
            ligne = {"code": code, "email": email,
                     "nom": emails_noms.get(email, "") or perf["vendeurs"].get(code, code),
                     "bilan": b, "objectifs": objs,
                     "rem_source": rem_source}
            (lignes if email else sans_collab).append(ligne)
        lignes.sort(key=lambda x: -(x["bilan"].get("ca_ttc_global") or 0))

    # Cumul des 12 derniers mois importés, par collaborateur associé
    cumul = []
    derniers = sorted(perf["mois"])[-12:]
    for code, email in perf["mapping"].items():
        if not email or email not in emails_noms:
            continue   # non associé, ou associé à un archivé : ignoré
        blocs = [perf["mois"][m][code] for m in derniers if code in perf["mois"][m]]
        if not blocs:
            continue
        total = sommer_mois(blocs)
        rem, _src = remuneration_effective(email, perf, profils)
        b = bilan_collaborateur(total, rem, options)
        cout_mensuel = b.pop("cout", None)
        b["cout"] = round(cout_mensuel * len(blocs), 2) if cout_mensuel is not None else None
        b["rentabilite"] = (round(b["marge_estimee"] - b["cout"], 2)
                            if b["marge_estimee"] is not None and b["cout"] is not None else None)
        b["ratio"] = (round(b["ca_ttc_global"] / b["cout"], 2)
                      if b.get("ca_ttc_global") and b.get("cout") else None)
        cumul.append({"email": email, "nom": emails_noms.get(email, code),
                      "nb_mois": len(blocs), "bilan": b})
    cumul.sort(key=lambda x: -(x["bilan"].get("rentabilite") if x["bilan"].get("rentabilite") is not None else -1e12))

    return render_template("admin_performance.html", vue="tableau",
                           perf=perf, options=options, mois_dispo=mois_dispo,
                           mois_sel=mois_sel, titre_periode=titre_periode,
                           lignes=lignes, sans_collab=sans_collab, cumul=cumul,
                           employes=employes, emails_noms=emails_noms,
                           indicateurs_obj=INDICATEURS_OBJECTIF,
                           crypto_ok=crypto_rh.crypto_disponible(),
                           msg=request.args.get("msg", ""))


def _page_collaborateur(perf, options, email, emails_noms, mois_dispo):
    """Vue détaillée : historique mensuel + rémunération + objectifs."""
    code = next((c for c, e in perf["mapping"].items() if e == email), None)
    rem_eff, rem_source = remuneration_effective(email, perf, charger_profils())
    histo = []
    for m in sorted(perf["mois"], reverse=True):
        bloc = perf["mois"][m].get(code or "")
        if bloc:
            histo.append({"mois": m,
                          "bilan": bilan_collaborateur(bloc, rem_eff, options)})
    rem_aff = {"brut": crypto_rh.dechiffrer(rem_eff["brut"]) if rem_eff.get("brut") else "",
               "variable": crypto_rh.dechiffrer(rem_eff["variable"]) if rem_eff.get("variable") else "",
               "source": rem_source}
    objectifs = [suivi_objectif({**o, "email": email}, perf, options)
                 for o in perf["objectifs"].get(email, [])]
    return render_template("admin_performance.html", vue="collab",
                           perf=perf, options=options, email=email, code=code,
                           nom=emails_noms.get(email, email), histo=histo,
                           rem=rem_aff, objectifs=objectifs,
                           mois_dispo=mois_dispo, titre_periode=titre_periode,
                           indicateurs_obj=INDICATEURS_OBJECTIF,
                           crypto_ok=crypto_rh.crypto_disponible(),
                           msg=request.args.get("msg", ""))


@bp.route("/admin/performance/import", methods=["POST"])
def importer():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    fichier = request.files.get("fichier")
    if not fichier or not fichier.filename.lower().endswith(".xlsx"):
        return redirect(url_for("performance.admin_performance", msg="fichier_ko"))
    vendeurs, mois, erreurs = importer_xlsx(fichier)
    if not mois:
        return redirect(url_for("performance.admin_performance", msg="import_ko"))
    perf = charger_perf()
    perf["vendeurs"].update(vendeurs)
    for m, bloc in mois.items():      # ré-importer un mois déjà connu l'écrase
        perf["mois"][m] = bloc
    profils_i = charger_profils()
    proposer_mapping(perf["vendeurs"], perf["mapping"],
                     [e for e in charger_employes()
                      if collaborateur_actif(profils_i.get(e["email"], {}))])
    perf["dernier_import"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance",
                            msg=f"import_ok:{len(mois)}"))


@bp.route("/admin/performance/mapping", methods=["POST"])
def enregistrer_mapping():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    perf = charger_perf()
    profils_m = charger_profils()
    emails_ok = {e["email"] for e in charger_employes()
                 if collaborateur_actif(profils_m.get(e["email"], {}))}
    for code in perf["vendeurs"]:
        champ = f"map_{code.replace(' ', '_')}"
        if champ in request.form:
            email = request.form.get(champ, "")
            perf["mapping"][code] = email if email in emails_ok else ""
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", msg="mapping_ok"))


@bp.route("/admin/performance/options", methods=["POST"])
def enregistrer_options():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    perf = charger_perf()
    for cle in OPTIONS_DEFAUT:
        try:
            perf["options"][cle] = float(request.form.get(cle, "").replace(",", "."))
        except ValueError:
            pass
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", msg="options_ok"))


@bp.route("/admin/performance/remuneration", methods=["POST"])
def enregistrer_remuneration():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    email = request.form.get("email", "")
    perf = charger_perf()
    rem = {}
    for champ in ("brut", "variable"):
        brut_txt = request.form.get(champ, "").strip().replace(",", ".")
        if brut_txt:
            try:
                rem[champ] = crypto_rh.chiffrer(str(float(brut_txt)))
            except ValueError:
                pass
    if rem:
        perf["remunerations"][email] = rem
    else:
        perf["remunerations"].pop(email, None)
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", collab=email, msg="rem_ok"))


@bp.route("/admin/performance/objectif", methods=["POST"])
def ajouter_objectif():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    email = request.form.get("email", "")
    indicateur = request.form.get("indicateur", "")
    if indicateur not in INDICATEURS_OBJECTIF:
        abort(400)
    try:
        cible = float(request.form.get("cible", "").replace(",", "."))
    except ValueError:
        return redirect(url_for("performance.admin_performance", collab=email, msg="cible_ko"))
    perf = charger_perf()
    perf["objectifs"].setdefault(email, []).append({
        "id": uuid.uuid4().hex[:8],
        "indicateur": indicateur,
        "cible": cible,
        "echeance": request.form.get("echeance", "").strip(),       # AAAA-MM
        "entretien": request.form.get("entretien", "").strip(),     # date de l'entretien
        "commentaire": request.form.get("commentaire", "").strip(),
        "clos": False,
    })
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", collab=email, msg="obj_ok"))


@bp.route("/admin/performance/objectif/clore", methods=["POST"])
def clore_objectif():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    email, oid = request.form.get("email", ""), request.form.get("id", "")
    perf = charger_perf()
    for o in perf["objectifs"].get(email, []):
        if o["id"] == oid:
            o["clos"] = not o.get("clos")
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", collab=email, msg="obj_maj"))


@bp.route("/admin/performance/objectif/supprimer", methods=["POST"])
def supprimer_objectif():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    email, oid = request.form.get("email", ""), request.form.get("id", "")
    perf = charger_perf()
    perf["objectifs"][email] = [o for o in perf["objectifs"].get(email, [])
                                if o["id"] != oid]
    sauvegarder_perf(perf)
    return redirect(url_for("performance.admin_performance", collab=email, msg="obj_suppr"))
